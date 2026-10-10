"""Bounded spreadsheet parsing and durable, explicitly applied catalog imports."""

import csv
import io
import json
import logging
import re
import zipfile
from collections import Counter
from datetime import timedelta
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import and_, or_, select

from salma.auth import rate_limit, require_business
from salma.db import now, session
from salma.platform_models import Business
from salma.product_import_models import ProductImport
from salma.settings import settings

router = APIRouter(prefix="/api/businesses/{business_id}")
log = logging.getLogger(__name__)
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_CELL = 8_000
FIELDS = {
    "sku",
    "name",
    "description",
    "category",
    "price",
    "currency",
    "stock_quantity",
    "stock_status",
    "conditions",
    "url",
    "options",
    "color",
    "size",
    "options.color",
    "options.size",
    "external_product_id",
    "external_variant_id",
}


def cell_text(value):
    text = "" if value is None else str(value).strip()
    if len(text) > MAX_CELL:
        raise ValueError(f"Cell text exceeds {MAX_CELL} characters")
    return text


def sheet_data(name, records):
    """Normalize headers and rows without evaluating spreadsheet expressions."""
    iterator = iter(records)
    try:
        header_values, header_formulas = next(iterator)
    except StopIteration:
        raise ValueError(f"Worksheet {name!r} is empty; add a header row") from None
    columns = [cell_text(v) for v in header_values]
    while columns and not columns[-1]:
        columns.pop()
    if not columns or any(not c for c in columns) or any(header_formulas):
        raise ValueError(f"Worksheet {name!r} needs a nonempty text header for each column")
    if len(columns) > MAX_COLUMNS:
        raise ValueError(f"Worksheet {name!r} exceeds {MAX_COLUMNS} columns")
    if len(set(columns)) != len(columns):
        raise ValueError(f"Worksheet {name!r} has duplicate headers; rename them")
    rows = []
    for number, (values, formulas) in enumerate(iterator, start=2):
        if number > MAX_ROWS + 1:
            raise ValueError(f"Worksheet {name!r} exceeds {MAX_ROWS} data rows")
        if len(values) > MAX_COLUMNS:
            raise ValueError(f"Row {number} exceeds {MAX_COLUMNS} columns")
        texts = [cell_text(v) for v in values]
        if not any(texts):
            continue
        if any(texts[len(columns) :]):
            raise ValueError(f"Row {number} has values in columns without headers")
        texts += [""] * (len(columns) - len(texts))
        rows.append(
            {
                "row_number": number,
                "values": texts[: len(columns)],
                "formulas": [i for i, formula in enumerate(formulas) if formula],
            }
        )
    return {"name": name, "columns": columns, "rows": rows}


def parse_file(path):
    path = Path(path)
    if path.stat().st_size > MAX_UPLOAD_BYTES:
        raise ValueError("Upload must be at most 20 MB")
    if path.suffix.lower() == ".csv":
        try:
            content = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError(
                "CSV must use UTF-8 encoding (Arabic and English are supported)"
            ) from None
        csv.field_size_limit(MAX_CELL)
        try:
            records = ((row, [False] * len(row)) for row in csv.reader(io.StringIO(content)))
            return [sheet_data("CSV", records)]
        except csv.Error as exc:
            raise ValueError(f"Invalid CSV: {exc}") from None
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Upload an XLSX or UTF-8 CSV file")
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 1_000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                raise ValueError("XLSX expands beyond the safe 50 MB limit")
            for entry in entries:
                if entry.flag_bits & 1:
                    raise ValueError("Encrypted XLSX files are not supported")
                # Check actual streams as well as untrusted ZIP directory sizes.
                with archive.open(entry) as stream:
                    total = 0
                    while block := stream.read(1024 * 1024):
                        total += len(block)
                        if total > MAX_EXPANDED_BYTES:
                            raise ValueError("XLSX expands beyond the safe 50 MB limit")
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) > 30:
                raise ValueError("XLSX may contain at most 30 worksheets")
            sheets = []
            row_count = 0
            for sheet in book.worksheets:
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise ValueError(f"Worksheet {sheet.title!r} exceeds {MAX_COLUMNS} columns")
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError(f"Worksheet {sheet.title!r} exceeds {MAX_ROWS} data rows")
                if sheet.max_row == 1 and sheet.max_column == 1 and sheet["A1"].value is None:
                    continue
                records = (
                    ([c.value for c in row], [c.data_type == "f" for c in row])
                    for row in sheet.iter_rows()
                )
                parsed = sheet_data(sheet.title, records)
                row_count += len(parsed["rows"])
                if row_count > MAX_ROWS:
                    raise ValueError(f"Workbook exceeds {MAX_ROWS} data rows across worksheets")
                sheets.append(parsed)
            if not sheets:
                raise ValueError("Workbook contains no headers or data")
            return sheets
        finally:
            book.close()
    except ValueError:
        raise
    except Exception:
        raise ValueError("Invalid or unsupported XLSX workbook; save as XLSX and retry") from None


def suspicious_csv(value, field):
    text = value.lstrip()
    if not text or text[0] not in "=+@-":
        return False
    # A negative number is numeric input, and the product validator rejects negative prices.
    return not (field in {"price", "stock_quantity"} and re.fullmatch(r"-\d+(\.\d+)?", text))


def preview_rows(item, existing, validator):
    details = dict(item.details)
    sheet = next(s for s in details["sheets"] if s["name"] == details["worksheet"])
    if not sheet["rows"]:
        raise ValueError("The selected worksheet has headers but no product rows")
    mapping = details["mapping"]
    indexes = {field: sheet["columns"].index(header) for field, header in mapping.items()}
    generated = dict(details.get("generated_skus", {}))
    existing_by_sku = {p["sku"].casefold(): p for p in existing}
    results = []
    for source in sheet["rows"]:
        number = source["row_number"]
        data = {field: source["values"][index] for field, index in indexes.items()}
        errors = []
        for field, index in indexes.items():
            value = source["values"][index]
            if index in source["formulas"] or (
                details["file_type"] == "csv" and suspicious_csv(value, field)
            ):
                errors.append(f"{field}: formulas are not accepted; replace with a literal value")
        key = f"{sheet['name']}:{number}"
        if not data.get("sku"):
            generated.setdefault(key, "IMP-" + uuid4().hex[:20].upper())
            data["sku"] = generated[key]
        try:
            options = json.loads(data.pop("options", "") or "{}")
            if not isinstance(options, dict):
                raise ValueError("options must be a JSON object")
            for field in ("color", "size", "options.color", "options.size"):
                value = data.pop(field, "")
                if value:
                    options[field.split(".")[-1]] = value
            data["options"] = options
            if not data.get("stock_quantity"):
                data["stock_quantity"] = None
            data.update(
                source_type="file",
                snapshot_at=details["snapshot_at"],
                source_ref=f"{item.filename} / {sheet['name']} / row {number}",
                archived=False,
            )
            canonical = validator(data, default_currency=details.get("default_currency"))
        except (ValueError, TypeError, json.JSONDecodeError) as exc:
            errors.append(str(exc))
            canonical = None
        result = {
            "row_number": number,
            "sku": data["sku"],
            "errors": errors,
            "product": canonical,
            "action": "invalid",
        }
        results.append(result)
    duplicates = Counter(row["sku"].casefold() for row in results)
    compare_ignored = {"source_type", "source_ref", "snapshot_at"}
    for result in results:
        if duplicates[result["sku"].casefold()] > 1:
            result["errors"].append(
                "sku: duplicate SKU in this file; each variant needs a unique SKU"
            )
        if not result["errors"]:
            old = existing_by_sku.get(result["sku"].casefold())
            product = result["product"]
            same = old is not None and all(
                old.get(k) == value for k, value in product.items() if k not in compare_ignored
            )
            result["action"] = "unchanged" if same else ("updated" if old else "new")
    merged = dict(existing_by_sku)
    merged.update({r["sku"].casefold(): r["product"] for r in results if not r["errors"]})
    if sum(not p.get("archived", False) for p in merged.values()) > MAX_ROWS:
        for result in results:
            result["errors"].append(
                "Catalog would exceed 10,000 active variants; archive products first"
            )
            result["action"] = "invalid"
    counts = {
        key: sum(r["action"] == key for r in results)
        for key in ("new", "updated", "unchanged", "invalid")
    }
    counts["total"] = len(results)
    details.update(rows=results, counts=counts, generated_skus=generated)
    return details


def process_product_import_job(import_id=None):
    """Claim DB work with a lease; interrupted jobs become claimable after five minutes."""
    token = str(uuid4())
    with session() as db:
        eligible = or_(
            ProductImport.status.in_(["pending", "pending_validation"]),
            and_(ProductImport.status == "processing", ProductImport.lease_until < now()),
        )
        query = select(ProductImport).where(eligible)
        if import_id:
            query = query.where(ProductImport.id == import_id)
        item = db.scalar(query.order_by(ProductImport.created_at).with_for_update(skip_locked=True))
        if item is None:
            return False
        item.status = "processing"
        item.lease_token = token
        item.lease_until = now() + timedelta(minutes=5)
        identifier = item.id
        stage = item.pending_stage
        path = item.storage_path
        db.commit()
    try:
        parsed = parse_file(path) if stage == "parse" else None
        with session() as db:
            item = db.scalar(
                select(ProductImport).where(ProductImport.id == identifier).with_for_update()
            )
            if item is None or item.lease_token != token:
                return True
            if stage == "parse":
                item.details = {**item.details, "sheets": parsed}
                item.status = "awaiting_mapping"
            else:
                from salma.product_catalog import get_draft_products, validate_product

                business = db.scalar(
                    select(Business).where(Business.id == item.business_id).with_for_update()
                )
                item.details = preview_rows(
                    item, get_draft_products(db, item.business_id), validate_product
                )
                item.preview_revision_id = business.draft_product_revision_id
                item.status = "invalid" if item.details["counts"]["invalid"] else "ready"
            item.error = None
            item.lease_token = None
            item.lease_until = None
            db.commit()
    except Exception as exc:
        log.exception("Product import %s failed", identifier)
        with session() as db:
            item = db.scalar(
                select(ProductImport).where(ProductImport.id == identifier).with_for_update()
            )
            if item and item.lease_token == token:
                item.status = "failed"
                item.error = (
                    str(exc)[:1000]
                    if isinstance(exc, ValueError)
                    else ("Import processing failed; check the file and retry the upload")
                )
                item.lease_token = None
                item.lease_until = None
                db.commit()
    return True


def import_row(db, business_id, import_id, lock=False):
    query = select(ProductImport).where(
        ProductImport.id == import_id, ProductImport.business_id == business_id
    )
    row = db.scalar(query.with_for_update() if lock else query)
    if row is None:
        raise HTTPException(404, "Product import not found")
    return row


def import_view(row, page=1, page_size=100, offset=None):
    details = row.details
    rows = details.get("rows", [])
    start = (page - 1) * page_size if offset is None else offset
    return {
        "id": row.id,
        "status": row.status,
        "filename": row.filename,
        "created_at": row.created_at.isoformat(),
        "error": row.error,
        "worksheets": [
            {"name": s["name"], "columns": s["columns"]} for s in details.get("sheets", [])
        ],
        "worksheet": details.get("worksheet"),
        "mapping": details.get("mapping", {}),
        "default_currency": details.get("default_currency"),
        "counts": details.get("counts", {}),
        "rows": rows[start : start + page_size],
        "total": len(rows),
        "offset": start,
        "limit": page_size,
        "pagination": {"page": page, "page_size": page_size, "total": len(rows)},
        "applied_revision_id": row.applied_revision_id,
    }


def csv_download(rows, filename):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    # Exported source references can contain user text; protect spreadsheet viewers too.
    for row in rows:
        writer.writerow(
            [("'" + str(v)) if str(v).lstrip().startswith(("=", "+", "-", "@")) else v for v in row]
        )
    return Response(
        "\ufeff" + output.getvalue(),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/product-imports/template")
def download_template(business_id: str, request: Request):
    require_business(request, business_id)
    return csv_download(
        [
            [
                "sku",
                "name",
                "description",
                "category",
                "color",
                "size",
                "price",
                "currency",
                "stock_quantity",
                "stock_status",
                "conditions",
                "url",
            ]
        ],
        "product-template.csv",
    )


@router.post("/product-imports", status_code=202)
async def upload_import(business_id: str, request: Request, file: UploadFile = File(...)):
    require_business(request, business_id)
    rate_limit("product-upload:" + business_id, 20)
    filename = (file.filename or "").replace("\\", "/").split("/")[-1]
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv"} or len(filename) > 255:
        raise HTTPException(
            422, "Upload an XLSX or UTF-8 CSV file with a filename under 256 characters"
        )
    identifier = str(uuid4())
    directory = Path(settings().uploads_dir) / business_id / "product-imports"
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / (identifier + suffix)
    try:
        size = 0
        with path.open("xb") as output:
            while block := await file.read(1024 * 1024):
                size += len(block)
                if size > MAX_UPLOAD_BYTES:
                    raise HTTPException(413, "Upload must be at most 20 MB")
                output.write(block)
        if not size:
            raise HTTPException(422, "The uploaded file is empty")
        with session() as db:
            item = ProductImport(
                id=identifier,
                business_id=business_id,
                filename=filename,
                storage_path=str(path.resolve()),
                status="pending",
                pending_stage="parse",
                details={
                    "file_type": suffix[1:],
                    "snapshot_at": now().isoformat(),
                    "generated_skus": {},
                },
            )
            db.add(item)
            db.commit()
            return import_view(item)
    except BaseException:
        path.unlink(missing_ok=True)
        raise
    finally:
        await file.close()


@router.get("/product-imports")
def list_imports(business_id: str, request: Request):
    require_business(request, business_id)
    with session() as db:
        return [
            import_view(row, page_size=0)
            for row in db.scalars(
                select(ProductImport)
                .where(ProductImport.business_id == business_id)
                .order_by(ProductImport.created_at.desc())
                .limit(100)
            )
        ]


@router.get("/product-imports/{import_id}")
def get_import(
    business_id: str,
    import_id: str,
    request: Request,
    page: int = Query(1, ge=1),
    page_size: int = Query(100, ge=1, le=500),
    offset: int | None = Query(None, ge=0),
    limit: int | None = Query(None, ge=1, le=500),
):
    require_business(request, business_id)
    with session() as db:
        return import_view(
            import_row(db, business_id, import_id),
            page,
            limit if limit is not None else page_size,
            offset,
        )


class ImportMapping(BaseModel):
    worksheet: str = Field(max_length=255)
    mapping: dict[str, str]
    default_currency: str | None = Field(default=None, max_length=3)


@router.put("/product-imports/{import_id}/mapping", status_code=202)
def save_mapping(business_id: str, import_id: str, body: ImportMapping, request: Request):
    require_business(request, business_id)
    with session() as db:
        row = import_row(db, business_id, import_id, lock=True)
        if row.status not in {"awaiting_mapping", "ready", "invalid", "failed"}:
            raise HTTPException(
                409, "Wait for parsing to finish; applied imports cannot be changed"
            )
        sheet = next(
            (s for s in row.details.get("sheets", []) if s["name"] == body.worksheet), None
        )
        if sheet is None:
            raise HTTPException(422, "Select a worksheet from this file")
        if set(body.mapping) - FIELDS or any(
            v not in sheet["columns"] for v in body.mapping.values()
        ):
            raise HTTPException(422, "Map supported product fields to existing column headers")
        if not {"name", "price"} <= body.mapping.keys():
            raise HTTPException(422, "Map at least name and price columns")
        currency = body.default_currency.upper() if body.default_currency else None
        if currency and not re.fullmatch("[A-Z]{3}", currency):
            raise HTTPException(422, "Default currency must be a three-letter currency code")
        if "currency" not in body.mapping and not currency:
            raise HTTPException(
                422, "Explicitly confirm a default currency or map a currency column"
            )
        row.details = {
            **row.details,
            "worksheet": body.worksheet,
            "mapping": body.mapping,
            "default_currency": currency,
            "rows": [],
            "counts": {},
        }
        row.status = "pending_validation"
        row.pending_stage = "validate"
        row.error = None
        db.commit()
        return import_view(row)


def apply_import(db, business_id, import_id):
    from salma.product_catalog import upsert_draft_products

    item = import_row(db, business_id, import_id, lock=True)
    if item.status == "applied":
        return item
    if item.status != "ready" or item.details.get("counts", {}).get("invalid", 1):
        raise HTTPException(409, "Validate the mapping and fix invalid rows before applying")
    business = db.scalar(select(Business).where(Business.id == business_id).with_for_update())
    if business.draft_product_revision_id != item.preview_revision_id:
        raise HTTPException(
            409, "Catalog changed after preview; save the mapping again to revalidate"
        )
    products = [r["product"] for r in item.details["rows"]]
    revision = business.draft_product_revision_id
    if products:
        revision = upsert_draft_products(
            db, business_id, products, expected_revision=item.preview_revision_id
        )
    item.status = "applied"
    item.applied_revision_id = revision
    return item


@router.post("/product-imports/{import_id}/apply")
def apply(business_id: str, import_id: str, request: Request):
    require_business(request, business_id)
    with session() as db:
        item = apply_import(db, business_id, import_id)
        db.commit()
        return import_view(item)


@router.get("/product-imports/{import_id}/generated-skus")
def generated_skus(business_id: str, import_id: str, request: Request):
    require_business(request, business_id)
    with session() as db:
        item = import_row(db, business_id, import_id)
        if item.status not in {"ready", "invalid", "applied"}:
            raise HTTPException(409, "Finish validation before downloading generated SKUs")
        details = item.details
        sheet = next(
            (s for s in details.get("sheets", []) if s["name"] == details.get("worksheet")), None
        )
        if sheet is None:
            raise HTTPException(
                409, "Choose a mapping and validate before downloading generated SKUs"
            )
        mapping = details.get("mapping", {})
        # Preserve original literal values for a reusable next import, adding generated identifiers.
        sku_header = mapping.get("sku", "sku")
        columns = list(sheet["columns"])
        if sku_header not in columns:
            columns.append(sku_header)
        index = columns.index(sku_header)
        rows = [columns]
        preview = {row["row_number"]: row["sku"] for row in details.get("rows", [])}
        for source in sheet["rows"]:
            values = list(source["values"]) + [""] * (len(columns) - len(source["values"]))
            if source["row_number"] in preview:
                values[index] = preview[source["row_number"]]
            rows.append(values)
        return csv_download(rows, "generated-skus.csv")
